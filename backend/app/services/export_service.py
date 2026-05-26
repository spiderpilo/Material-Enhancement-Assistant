from io import BytesIO

from pptx import Presentation
from pptx.util import Pt

from app.models.slide_deck_model import SlideDeckOutline


class SlideDeckExportError(Exception):
    """Raised when a generated slide deck cannot be exported."""


def build_slide_deck_pptx_bytes(*, outline: SlideDeckOutline) -> bytes:
    if not outline.slides:
        raise SlideDeckExportError("Cannot export a slide deck without slides.")

    presentation = Presentation()

    title_slide_layout = presentation.slide_layouts[0]
    title_slide = presentation.slides.add_slide(title_slide_layout)
    title_shape = title_slide.shapes.title
    subtitle_shape = title_slide.placeholders[1] if len(title_slide.placeholders) > 1 else None

    if title_shape is not None:
        title_shape.text = outline.title.strip() or "Generated Slide Deck"
    if subtitle_shape is not None:
        subtitle_shape.text = (outline.subtitle or "").strip()

    for slide_outline in outline.slides:
        body_layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(body_layout)
        title = slide.shapes.title
        body_placeholder = slide.shapes.placeholders[1] if len(slide.shapes.placeholders) > 1 else None

        if title is not None:
            title.text = slide_outline.title.strip() or "Untitled Slide"

        if body_placeholder is None:
            continue

        text_frame = body_placeholder.text_frame
        text_frame.clear()
        text_frame.word_wrap = True

        bullets = [bullet.strip() for bullet in slide_outline.bullets if bullet.strip()]
        if not bullets:
            bullets = ["Key points unavailable for this slide."]

        for index, bullet in enumerate(bullets):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(24)

        notes_text = slide_outline.speaker_notes.strip()
        if notes_text:
            notes_frame = slide.notes_slide.notes_text_frame
            notes_frame.clear()
            notes_frame.text = notes_text

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
